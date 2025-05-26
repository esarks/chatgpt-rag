from flask import Flask, request, jsonify, Response
from flask_cors import CORS
import os
from openai import OpenAI
from pinecone import Pinecone
from dotenv import load_dotenv
import tempfile
import json
import re
import openpyxl
import pandas as pd
import time

load_dotenv()

app = Flask(__name__)
CORS(app)

client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
pc = Pinecone(api_key=os.getenv("PINECONE_API_KEY"))
index = pc.Index(os.getenv("PINECONE_INDEX"))

FILE_LOG = "uploaded_files.json"
CHUNK_LOG = "chunk_log.json"

def sanitize_id(text):
    return re.sub(r'[^A-Za-z0-9_.-]', '_', text)

def extract_text(file_path, file_type):
    text = ""
    print(f"🔍 Extracting: {file_path} ({file_type})")

    if file_type.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text += page.extract_text() or ""

    elif file_type.endswith(".docx"):
        import docx
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"

    elif file_type.endswith(".txt"):
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()

    elif file_type.endswith(".xlsx"):
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) if cell is not None else "" for cell in row])
                    text += row_text + "\n"
        except Exception as e:
            print(f"❌ openpyxl failed: {e} – trying pandas")
            try:
                df = pd.read_excel(file_path, engine="openpyxl")
                text = df.astype(str).apply(" ".join, axis=1).str.cat(sep="\n")
            except Exception as e2:
                raise ValueError(f"Unable to read Excel file using openpyxl or pandas: {e2}")

    else:
        raise ValueError("Unsupported file type")
    return text

@app.route("/upload", methods=["POST"])
def upload_file():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    file = request.files["file"]
    if file.filename == "":
        return jsonify({"error": "Empty filename"}), 400

    try:
        print(f"🔁 Starting upload for {file.filename}")
        start_time = time.time()

        with tempfile.NamedTemporaryFile(delete=False) as temp_file:
            file.save(temp_file.name)
            print(f"✅ File saved at {temp_file.name}")
            text = extract_text(temp_file.name, file.filename.lower())

        if not text.strip():
            raise ValueError("File is empty or failed to extract any text.")

        chunks = [text[i:i+500] for i in range(0, len(text), 500)]
        print(f"🧩 Created {len(chunks)} chunks")

        for i, chunk in enumerate(chunks):
            embedding = client.embeddings.create(
                model="text-embedding-ada-002",
                input=chunk
            ).data[0].embedding

            vector_id = f"{sanitize_id(file.filename)}-{i}"
            index.upsert([(vector_id, embedding, {"text": chunk, "source": file.filename})])

        uploaded = set()
        if os.path.exists(FILE_LOG):
            with open(FILE_LOG, "r") as f:
                uploaded = set(json.load(f))
        uploaded.add(file.filename)
        with open(FILE_LOG, "w") as f:
            json.dump(sorted(uploaded), f)

        chunk_map = {}
        if os.path.exists(CHUNK_LOG):
            with open(CHUNK_LOG, "r") as f:
                chunk_map = json.load(f)
        chunk_map[file.filename] = len(chunks)
        with open(CHUNK_LOG, "w") as f:
            json.dump(chunk_map, f)

        print(f"✅ Finished processing {file.filename} in {round(time.time() - start_time, 2)}s")
        return jsonify({
            "message": f"Uploaded and indexed {len(chunks)} chunks from {file.filename}.",
            "filename": file.filename,
            "chunks": len(chunks)
        })

    except Exception as e:
        print(f"❌ Error during upload of {file.filename}: {e}")
        return jsonify({"error": str(e)}), 500

@app.route("/files", methods=["GET"])
def list_uploaded_files():
    try:
        if os.path.exists(FILE_LOG):
            with open(FILE_LOG, "r") as f:
                uploaded = json.load(f)
        else:
            uploaded = []
        return jsonify({"files": sorted(uploaded)})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/files/details", methods=["GET"])
def file_chunk_details():
    try:
        with open(CHUNK_LOG, "r") as f:
            chunk_map = json.load(f)
        result = [{"filename": k, "chunks": v} for k, v in sorted(chunk_map.items())]
        return jsonify(result)
    except Exception:
        return jsonify([])

@app.route("/ask", methods=["POST"])
def ask_question():
    data = request.get_json()
    question = data.get("question", "")
    if not question:
        return jsonify({"error": "No question provided"}), 400

    embedding = client.embeddings.create(
        model="text-embedding-ada-002",
        input=question
    ).data[0].embedding

    pinecone_response = index.query(vector=embedding, top_k=5, include_metadata=True)

    context = ""
    sources = []
    for match in pinecone_response["matches"]:
        metadata = match.get("metadata", {})
        context += metadata.get("text", "") + "\n"
        sources.append(metadata.get("source", "unknown"))

    chat_response = client.chat.completions.create(
        model="gpt-3.5-turbo",
        messages=[
            {"role": "system", "content": "You are a helpful assistant that answers questions using the provided context."},
            {"role": "user", "content": f"Context:\n{context}\n\nQuestion: {question}"}
        ]
    )

    return jsonify({
        "question": question,
        "answer": chat_response.choices[0].message.content,
        "sources": list(set(sources))
    })

@app.route("/stream", methods=["POST"])
def stream_answer():
    data = request.get_json()
    question = data.get("question", "")
    if not question:
        return jsonify({"error": "No question provided"}), 400

    embedding = client.embeddings.create(
        model="text-embedding-ada-002",
        input=question
    ).data[0].embedding

    pinecone_response = index.query(vector=embedding, top_k=5, include_metadata=True)

    context = ""
    for match in pinecone_response["matches"]:
        metadata = match.get("metadata", {})
        context += metadata.get("text", "") + "\n"

    def generate():
        chat_stream = client.chat.completions.create(
            model="gpt-3.5-turbo",
            stream=True,
            messages=[
                {"role": "system", "content": "You are a helpful assistant that answers questions using the provided context."},
                {"role": "user", "content": f"Context:\n{context}\n\nQuestion: {question}"}
            ]
        )
        for chunk in chat_stream:
            if chunk.choices and chunk.choices[0].delta.content:
                yield chunk.choices[0].delta.content

    return Response(generate(), mimetype="text/plain")

if __name__ == "__main__":
    import os
    port = int(os.environ.get("PORT", 8080))
    app.run(host="0.0.0.0", port=port)